"""
DOCX Parser Module
Handles Microsoft Word document parsing and text extraction
"""

import os
from typing import Optional
from docx import Document
from .base_parser import BaseParser, ParseResult, DocumentMetadata


class DOCXParser(BaseParser):
    """Parser for DOCX documents"""
    
    def __init__(self, language: str = 'en', preserve_diacritics: bool = True):
        """
        Initialize DOCX parser
        
        Args:
            language: Document language
            preserve_diacritics: Whether to preserve Arabic diacritics
        """
        super().__init__(language, preserve_diacritics)
    
    def parse(self, file_path: str) -> ParseResult:
        """
        Parse DOCX document
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            ParseResult with extracted content
        """
        try:
            # Validate file
            if not os.path.exists(file_path):
                return ParseResult(
                    document_id='',
                    text='',
                    chunks=[],
                    metadata=None,
                    success=False,
                    error=f"File not found: {file_path}"
                )
            
            # Generate document ID
            document_id = self._generate_document_id(file_path)
            
            # Extract text
            text = self.extract_text(file_path)
            
            # Clean text
            text = self._clean_text(text)
            
            # Calculate statistics
            stats = self._calculate_stats(text)
            
            # Get file metadata
            file_size = os.path.getsize(file_path)
            filename = os.path.basename(file_path)
            
            # Count pages (approximate for DOCX)
            num_pages = self._estimate_pages(text)
            
            # Detect language
            detected_language = 'ar' if self.is_arabic(text) else self.language
            
            # Create metadata
            metadata = DocumentMetadata(
                document_id=document_id,
                filename=filename,
                file_type='docx',
                file_size=file_size,
                num_pages=num_pages,
                language=detected_language,
                total_chars=stats['total_chars'],
                total_words=stats['total_words'],
                metadata={'parser_version': '1.0.0'}
            )
            
            # Create result
            result = ParseResult(
                document_id=document_id,
                text=text,
                chunks=[],
                metadata=metadata,
                success=True
            )
            
            return result
            
        except Exception as e:
            return ParseResult(
                document_id='',
                text='',
                chunks=[],
                metadata=None,
                success=False,
                error=f"DOCX parsing error: {str(e)}"
            )
    
    def extract_text(self, file_path: str) -> str:
        """
        Extract raw text from DOCX
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Extracted text
        """
        doc = Document(file_path)
        text_parts = []
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        
        return '\n'.join(text_parts)
    
    def _estimate_pages(self, text: str, chars_per_page: int = 3000) -> int:
        """
        Estimate number of pages based on character count
        
        Args:
            text: Document text
            chars_per_page: Average characters per page
            
        Returns:
            Estimated page count
        """
        return max(1, len(text) // chars_per_page)
    
    def extract_metadata(self, file_path: str) -> dict:
        """
        Extract DOCX metadata
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Dictionary with DOCX metadata
        """
        metadata = {}
        
        try:
            doc = Document(file_path)
            core_properties = doc.core_properties
            
            metadata = {
                'title': core_properties.title or '',
                'author': core_properties.author or '',
                'subject': core_properties.subject or '',
                'keywords': core_properties.keywords or '',
                'created': str(core_properties.created) if core_properties.created else '',
                'modified': str(core_properties.modified) if core_properties.modified else '',
                'revision': core_properties.revision or 0
            }
        except Exception as e:
            print(f"Warning: Could not extract DOCX metadata: {e}")
        
        return metadata
